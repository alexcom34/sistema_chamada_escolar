import tkinter as tk
from tkinter import ttk, messagebox, filedialog

from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph

from repositories.relatorio_presenca_repository import (
    listar_turmas_disciplinas_relatorio,
    listar_alunos_relatorio,
    gerar_relatorio_presenca
)


class RelatorioPresencaView:

    def __init__(self, parent):

        self.janela = tk.Toplevel(parent)
        self.janela.title("Relatório Detalhado de Presença")
        self.janela.geometry("1200x700")
        self.janela.minsize(1000, 600)

        self.opcoes = []
        self.alunos = []

        self.dados_relatorio = []
        self.aulas_relatorio = {}

        self.tree_relatorio = None

        self.modo = tk.StringVar(value="turma")

        self.criar_interface()
        self.carregar_turmas_disciplinas()

    # =========================================================
    # INTERFACE
    # =========================================================

    def criar_interface(self):

        frame_filtros = ttk.LabelFrame(
            self.janela,
            text="Filtros do relatório"
        )

        frame_filtros.pack(
            fill="x",
            padx=10,
            pady=10
        )

        # -----------------------------------------------------
        # MODO
        # -----------------------------------------------------

        ttk.Label(
            frame_filtros,
            text="Exibir:"
        ).grid(
            row=0,
            column=0,
            padx=5,
            pady=5,
            sticky="w"
        )

        ttk.Radiobutton(
            frame_filtros,
            text="Turma inteira",
            variable=self.modo,
            value="turma",
            command=self.alterar_modo
        ).grid(
            row=0,
            column=1,
            padx=5,
            pady=5
        )

        ttk.Radiobutton(
            frame_filtros,
            text="Aluno",
            variable=self.modo,
            value="aluno",
            command=self.alterar_modo
        ).grid(
            row=0,
            column=2,
            padx=5,
            pady=5
        )

        # -----------------------------------------------------
        # TURMA / DISCIPLINA
        # -----------------------------------------------------

        ttk.Label(
            frame_filtros,
            text="Turma / Disciplina:"
        ).grid(
            row=1,
            column=0,
            padx=5,
            pady=5,
            sticky="w"
        )

        self.combo_turma_disciplina = ttk.Combobox(
            frame_filtros,
            width=55,
            state="readonly"
        )

        self.combo_turma_disciplina.grid(
            row=1,
            column=1,
            columnspan=2,
            padx=5,
            pady=5,
            sticky="ew"
        )

        self.combo_turma_disciplina.bind(
            "<<ComboboxSelected>>",
            self.carregar_alunos
        )

        # -----------------------------------------------------
        # ALUNO
        # -----------------------------------------------------

        ttk.Label(
            frame_filtros,
            text="Aluno:"
        ).grid(
            row=2,
            column=0,
            padx=5,
            pady=5,
            sticky="w"
        )

        self.combo_aluno = ttk.Combobox(
            frame_filtros,
            width=55,
            state="disabled"
        )

        self.combo_aluno.grid(
            row=2,
            column=1,
            columnspan=2,
            padx=5,
            pady=5,
            sticky="ew"
        )

        # -----------------------------------------------------
        # BOTÕES
        # -----------------------------------------------------

        frame_botoes = ttk.Frame(frame_filtros)

        frame_botoes.grid(
            row=3,
            column=0,
            columnspan=3,
            pady=10
        )

        ttk.Button(
            frame_botoes,
            text="Gerar Relatório",
            command=self.gerar
        ).pack(
            side="left",
            padx=5
        )

        ttk.Button(
            frame_botoes,
            text="Exportar Excel",
            command=self.exportar_excel
        ).pack(
            side="left",
            padx=5
        )

        ttk.Button(
            frame_botoes,
            text="Exportar PDF",
            command=self.exportar_pdf
        ).pack(
            side="left",
            padx=5
        )

        frame_filtros.columnconfigure(1, weight=1)

        # =====================================================
        # ÁREA DO RELATÓRIO
        # =====================================================

        frame_relatorio = ttk.Frame(self.janela)

        frame_relatorio.pack(
            fill="both",
            expand=True,
            padx=10,
            pady=(0, 10)
        )

        # Scroll vertical
        scrollbar_vertical = ttk.Scrollbar(
            frame_relatorio,
            orient="vertical"
        )

        scrollbar_vertical.pack(
            side="right",
            fill="y"
        )

        # Scroll horizontal
        scrollbar_horizontal = ttk.Scrollbar(
            frame_relatorio,
            orient="horizontal"
        )

        scrollbar_horizontal.pack(
            side="bottom",
            fill="x"
        )

        self.frame_tabela = ttk.Frame(frame_relatorio)

        self.frame_tabela.pack(
            fill="both",
            expand=True
        )

        # =====================================================
        # ESTILO DO CABEÇALHO
        # =====================================================

        style = ttk.Style()

        try:
            style.configure(
                "Relatorio.Treeview.Heading",
                padding=(5, 14),
                font=("Arial", 9, "bold")
            )
        except Exception:
            pass

        # =====================================================
        # TREEVIEW
        # =====================================================

        self.tree_relatorio = ttk.Treeview(
            self.frame_tabela,
            show="headings"
        )

        self.tree_relatorio.pack(
            fill="both",
            expand=True
        )

        self.tree_relatorio.configure(
            yscrollcommand=scrollbar_vertical.set,
            xscrollcommand=scrollbar_horizontal.set,
            style="Relatorio.Treeview"
        )

        scrollbar_vertical.configure(
            command=self.tree_relatorio.yview
        )

        scrollbar_horizontal.configure(
            command=self.tree_relatorio.xview
        )

    # =========================================================
    # CARREGAR TURMAS / DISCIPLINAS
    # =========================================================

    def carregar_turmas_disciplinas(self):

        self.opcoes = listar_turmas_disciplinas_relatorio()

        valores = []

        for item in self.opcoes:

            texto = (
                f"{item['turma']} - "
                f"{item['disciplina']}"
            )

            valores.append(texto)

        self.combo_turma_disciplina["values"] = valores

        if valores:
            self.combo_turma_disciplina.current(0)
            self.carregar_alunos()

    # =========================================================
    # ALTERAR MODO
    # =========================================================

    def alterar_modo(self):

        if self.modo.get() == "turma":

            self.combo_aluno.set("")

            self.combo_aluno.configure(
                state="disabled"
            )

        else:

            self.combo_aluno.configure(
                state="readonly"
            )

            self.carregar_alunos()

    # =========================================================
    # CARREGAR ALUNOS
    # =========================================================

    def carregar_alunos(self, event=None):

        indice = self.combo_turma_disciplina.current()

        if indice < 0:
            return

        if indice >= len(self.opcoes):
            return

        turma_id = self.opcoes[indice]["turma_id"]

        self.alunos = listar_alunos_relatorio(turma_id)

        valores = []

        for aluno in self.alunos:

            texto = (
                f"{aluno['nome']} "
                f"({aluno['matricula']})"
            )

            valores.append(texto)

        self.combo_aluno["values"] = valores

        if valores and self.modo.get() == "aluno":
            self.combo_aluno.current(0)

    # =========================================================
    # GERAR RELATÓRIO
    # =========================================================

    def gerar(self):

        indice = self.combo_turma_disciplina.current()

        if indice < 0:

            messagebox.showwarning(
                "Atenção",
                "Selecione uma turma e disciplina."
            )

            return

        turma_disciplina_id = self.opcoes[indice][
            "turma_disciplina_id"
        ]

        aluno_id = None

        if self.modo.get() == "aluno":

            indice_aluno = self.combo_aluno.current()

            if indice_aluno < 0:

                messagebox.showwarning(
                    "Atenção",
                    "Selecione um aluno."
                )

                return

            aluno_id = self.alunos[indice_aluno]["id"]

        try:

            dados = gerar_relatorio_presenca(
                turma_disciplina_id
            )

            self.mostrar_tabela(
                dados,
                aluno_id
            )

        except Exception as erro:

            messagebox.showerror(
                "Erro",
                f"Não foi possível gerar o relatório:\n\n{erro}"
            )

    # =========================================================
    # MONTAR TABELA
    # =========================================================

    def mostrar_tabela(self, dados, aluno_id=None):

        # Limpa tabela anterior

        for item in self.tree_relatorio.get_children():
            self.tree_relatorio.delete(item)

        self.tree_relatorio["columns"] = ()

        self.dados_relatorio = dados

        # -----------------------------------------------------
        # FILTRAR ALUNO
        # -----------------------------------------------------

        if aluno_id is not None:

            dados_filtrados = [
                item
                for item in dados
                if item["aluno_id"] == aluno_id
            ]

        else:

            dados_filtrados = dados

        if not dados_filtrados:

            messagebox.showinfo(
                "Relatório",
                "Não existem registros de presença para os filtros selecionados."
            )

            return

        # -----------------------------------------------------
        # DESCOBRIR TODAS AS AULAS
        # -----------------------------------------------------

        aulas = {}

        for item in dados_filtrados:

            chave = (
                item["aula_id"],
                item["data_aula"],
                item["numero_hora_aula"]
            )

            aulas[chave] = {
                "aula_id": item["aula_id"],
                "data_aula": item["data_aula"],
                "hora": item["numero_hora_aula"]
            }

        # Ordena por data e hora

        aulas_ordenadas = sorted(
            aulas.values(),
            key=lambda x: (
                x["data_aula"],
                x["hora"]
            )
        )

        self.aulas_relatorio = aulas_ordenadas

        # -----------------------------------------------------
        # COLUNAS
        # -----------------------------------------------------

        colunas = [
            "aluno",
            "matricula"
        ]

        for i, aula in enumerate(aulas_ordenadas):

            coluna = f"aula_{i}"

            colunas.append(coluna)

        self.tree_relatorio["columns"] = colunas

        # -----------------------------------------------------
        # COLUNAS FIXAS
        # -----------------------------------------------------

        self.tree_relatorio.heading(
            "aluno",
            text="Aluno"
        )

        self.tree_relatorio.heading(
            "matricula",
            text="Matrícula"
        )

        self.tree_relatorio.column(
            "aluno",
            width=220,
            minwidth=180,
            anchor="w"
        )

        self.tree_relatorio.column(
            "matricula",
            width=100,
            minwidth=90,
            anchor="center"
        )

        # -----------------------------------------------------
        # COLUNAS DAS AULAS
        # -----------------------------------------------------

        for i, aula in enumerate(aulas_ordenadas):

            coluna = f"aula_{i}"

            data = aula["data_aula"]

            if hasattr(data, "strftime"):
                data_formatada = data.strftime("%d/%m")
            else:
                data_formatada = str(data)

            # Cabeçalho em duas linhas
            titulo = (
                f"{data_formatada}\n"
                f"H{aula['hora']}"
            )

            self.tree_relatorio.heading(
                coluna,
                text=titulo
            )

            self.tree_relatorio.column(
                coluna,
                width=75,
                minwidth=75,
                anchor="center",
                stretch=False
            )

        # -----------------------------------------------------
        # MAPA DE PRESENÇAS
        # -----------------------------------------------------

        mapa = {}

        for item in dados_filtrados:

            chave = (
                item["aluno_id"],
                item["aula_id"],
                item["numero_hora_aula"]
            )

            mapa[chave] = item["status"]

        # -----------------------------------------------------
        # ALUNOS
        # -----------------------------------------------------

        alunos = {}

        for item in dados_filtrados:

            alunos[item["aluno_id"]] = {
                "nome": item["aluno"],
                "matricula": item["matricula"]
            }

        # -----------------------------------------------------
        # INSERE LINHAS
        # -----------------------------------------------------

        for aluno_id_atual, aluno in sorted(
            alunos.items(),
            key=lambda x: x[1]["nome"]
        ):

            valores = [
                aluno["nome"],
                aluno["matricula"]
            ]

            for aula in aulas_ordenadas:

                chave = (
                    aluno_id_atual,
                    aula["aula_id"],
                    aula["hora"]
                )

                status = mapa.get(chave)

                if status == "PRESENTE":
                    valor = "P"

                elif status == "FALTA":
                    valor = "F"

                elif status == "JUSTIFICADA":
                    valor = "J"

                else:
                    valor = "-"

                valores.append(valor)

            self.tree_relatorio.insert(
                "",
                "end",
                values=valores
            )

    # =========================================================
    # PREPARAR DADOS PARA EXPORTAÇÃO
    # =========================================================

    def obter_dados_exportacao(self):

        if not self.dados_relatorio:
            return None, None

        dados = self.dados_relatorio

        # Verifica se está filtrando um aluno

        aluno_id = None

        if self.modo.get() == "aluno":

            indice_aluno = self.combo_aluno.current()

            if indice_aluno >= 0:
                aluno_id = self.alunos[indice_aluno]["id"]

        if aluno_id is not None:

            dados = [
                item
                for item in dados
                if item["aluno_id"] == aluno_id
            ]

        if not dados:
            return None, None

        # -----------------------------------------------------
        # AULAS
        # -----------------------------------------------------

        aulas = {}

        for item in dados:

            chave = (
                item["aula_id"],
                item["data_aula"],
                item["numero_hora_aula"]
            )

            aulas[chave] = {
                "aula_id": item["aula_id"],
                "data_aula": item["data_aula"],
                "hora": item["numero_hora_aula"]
            }

        aulas_ordenadas = sorted(
            aulas.values(),
            key=lambda x: (
                x["data_aula"],
                x["hora"]
            )
        )

        # -----------------------------------------------------
        # MAPA
        # -----------------------------------------------------

        mapa = {}

        for item in dados:

            chave = (
                item["aluno_id"],
                item["aula_id"],
                item["numero_hora_aula"]
            )

            if item["status"] == "PRESENTE":
                valor = "P"

            elif item["status"] == "FALTA":
                valor = "F"

            elif item["status"] == "JUSTIFICADA":
                valor = "J"

            else:
                valor = "-"

            mapa[chave] = valor

        # -----------------------------------------------------
        # ALUNOS
        # -----------------------------------------------------

        alunos = {}

        for item in dados:

            alunos[item["aluno_id"]] = {
                "nome": item["aluno"],
                "matricula": item["matricula"]
            }

        linhas = []

        for aluno_id_atual, aluno in sorted(
            alunos.items(),
            key=lambda x: x[1]["nome"]
        ):

            linha = [
                aluno["nome"],
                aluno["matricula"]
            ]

            for aula in aulas_ordenadas:

                chave = (
                    aluno_id_atual,
                    aula["aula_id"],
                    aula["hora"]
                )

                linha.append(
                    mapa.get(chave, "-")
                )

            linhas.append(linha)

        # -----------------------------------------------------
        # CABEÇALHOS
        # -----------------------------------------------------

        cabecalhos = [
            "Aluno",
            "Matrícula"
        ]

        for aula in aulas_ordenadas:

            data = aula["data_aula"]

            if hasattr(data, "strftime"):
                data_formatada = data.strftime("%d/%m/%Y")
            else:
                data_formatada = str(data)

            cabecalhos.append(
                f"{data_formatada} H{aula['hora']}"
            )

        return cabecalhos, linhas

    # =========================================================
    # EXPORTAR EXCEL
    # =========================================================

    def exportar_excel(self):

        cabecalhos, linhas = self.obter_dados_exportacao()

        if not cabecalhos:

            messagebox.showwarning(
                "Exportação",
                "Gere um relatório antes de exportar."
            )

            return

        caminho = filedialog.asksaveasfilename(
            title="Salvar relatório em Excel",
            defaultextension=".xlsx",
            filetypes=[
                (
                    "Arquivo Excel",
                    "*.xlsx"
                )
            ]
        )

        if not caminho:
            return

        try:

            workbook = Workbook()

            planilha = workbook.active
            planilha.title = "Presença"

            # -------------------------------------------------
            # CABEÇALHO
            # -------------------------------------------------

            for coluna, valor in enumerate(
                cabecalhos,
                start=1
            ):

                celula = planilha.cell(
                    row=1,
                    column=coluna,
                    value=valor
                )

                celula.font = celula.font.copy(
                    bold=True
                )

            # -------------------------------------------------
            # DADOS
            # -------------------------------------------------

            for linha_excel, linha in enumerate(
                linhas,
                start=2
            ):

                for coluna_excel, valor in enumerate(
                    linha,
                    start=1
                ):

                    planilha.cell(
                        row=linha_excel,
                        column=coluna_excel,
                        value=valor
                    )

            # -------------------------------------------------
            # AJUSTE DE LARGURA
            # -------------------------------------------------

            for coluna in planilha.columns:

                letra = coluna[0].column_letter

                maior = 0

                for celula in coluna:

                    if celula.value is not None:

                        tamanho = len(
                            str(celula.value)
                        )

                        if tamanho > maior:
                            maior = tamanho

                planilha.column_dimensions[
                    letra
                ].width = min(
                    maior + 2,
                    30
                )

            # Congela Aluno/Matrícula

            planilha.freeze_panes = "C2"

            workbook.save(caminho)

            messagebox.showinfo(
                "Exportação concluída",
                f"Arquivo Excel salvo com sucesso:\n\n{caminho}"
            )

        except Exception as erro:

            messagebox.showerror(
                "Erro",
                f"Não foi possível exportar para Excel:\n\n{erro}"
            )

    # =========================================================
    # EXPORTAR PDF
    # =========================================================

    def exportar_pdf(self):

        cabecalhos, linhas = self.obter_dados_exportacao()

        if not cabecalhos:

            messagebox.showwarning(
                "Exportação",
                "Gere um relatório antes de exportar."
            )

            return

        caminho = filedialog.asksaveasfilename(
            title="Salvar relatório em PDF",
            defaultextension=".pdf",
            filetypes=[
                (
                    "Arquivo PDF",
                    "*.pdf"
                )
            ]
        )

        if not caminho:
            return

        try:

            documento = SimpleDocTemplate(
                caminho,
                pagesize=landscape(A4),
                rightMargin=20,
                leftMargin=20,
                topMargin=20,
                bottomMargin=20
            )

            estilos = getSampleStyleSheet()

            elementos = []

            # -------------------------------------------------
            # TÍTULO
            # -------------------------------------------------

            elementos.append(
                Paragraph(
                    "Relatório Detalhado de Presença",
                    estilos["Title"]
                )
            )

            elementos.append(
                Paragraph(
                    "P = Presente | F = Falta | J = Justificada | - = Sem aula",
                    estilos["Normal"]
                )
            )

            # -------------------------------------------------
            # TABELA
            # -------------------------------------------------

            tabela_dados = [
                cabecalhos
            ]

            tabela_dados.extend(
                linhas
            )

            tabela = Table(
                tabela_dados,
                repeatRows=1
            )

            tabela.setStyle(
                TableStyle([
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.lightgrey
                    ),

                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.black
                    ),

                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold"
                    ),

                    (
                        "FONTSIZE",
                        (0, 0),
                        (-1, -1),
                        7
                    ),

                    (
                        "ALIGN",
                        (1, 0),
                        (-1, -1),
                        "CENTER"
                    ),

                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "MIDDLE"
                    ),

                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.black
                    ),

                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        4
                    ),

                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        4
                    )
                ])
            )

            elementos.append(
                tabela
            )

            documento.build(
                elementos
            )

            messagebox.showinfo(
                "Exportação concluída",
                f"Arquivo PDF salvo com sucesso:\n\n{caminho}"
            )

        except Exception as erro:

            messagebox.showerror(
                "Erro",
                f"Não foi possível exportar para PDF:\n\n{erro}"
            )